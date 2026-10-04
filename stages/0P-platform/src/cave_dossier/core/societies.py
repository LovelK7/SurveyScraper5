"""Registar udruga: every caving society and the forms it goes by.

Two committed files beside this module, one registry:

* ``crospeleo_organizations.json`` — the **ground truth** (user, 2026-10-04):
  every organisation CroSpeleo has ever credited, under its canonical name,
  mined from a CroSpeleo objects export by ``cavedossier societies build``.
* ``societies.json`` — the curated overlay, joined on the canonical: the
  working name, a ``short`` form for tight spaces, and the spellings OSZs and
  SB actually carry (wrong ones included — SB 1328's OSZ calls the Bitelić
  Speleo *sekcija* an "SO"). It may also list an HPS society CroSpeleo has
  never credited; ``societies check`` reports those.

An organisation the overlay does not mention still gets a working name derived
from its canonical (``Speleološki odsjek HPD "Mosor", Split`` -> ``SO HPD
Mosor``) and, where the four-prefix rule gives one nobody else claims, a short
form (``SOM``).

``find`` resolves any written form to its society; ``fit_societies`` rewrites a
list into shorter forms, a step at a time, until it fits where it is printed.
Why a registry and not just ``people.society_shorthand``: the rule only knows
the ``SO/SD/SK/SU + named entity`` shapes. ``SKOL`` is already short and the
rule cannot see it as a society, and ``SO Sv. Jakov Bitelić`` comes out
``SOS`` — which nobody writes, for a society that is not even an odsjek. SB
1328's Istražili overflowed the Sastavnica for exactly that reason (user,
2026-10-04). The rule stays as the fallback for an unlisted society.

Lookup is exact on a diacritic-, case- and punctuation-insensitive key
(``normalize_lookup_key``), so ``Speleološka udruga "Estavela", Kastav`` and
``speleoloska udruga estavela kastav`` are one key; every form is also
accepted with the seat appended. Curated forms win over derived ones, and a
derived key two societies would share resolves to neither. Never fuzzy: a
wrong society on a printed nacrt is worse than an unabbreviated one.
"""

from __future__ import annotations

import json
import re
from collections import Counter
from dataclasses import dataclass, field
from datetime import date
from functools import lru_cache
from pathlib import Path
from typing import Callable

from cave_dossier.core.normalization import normalize_lookup_key
from cave_dossier.core.people import society_shorthand

OVERLAY_PATH = Path(__file__).parent / "societies.json"
CROSPELEO_PATH = Path(__file__).parent / "crospeleo_organizations.json"

# Type prefixes, long form -> the short form a working name uses.
_TYPE_PREFIXES: tuple[tuple[str, str], ...] = (
    ("Speleološko-alpinistički klub", "SAK"),
    ("Speleološki odsjek", "SO"),
    ("Speleološko društvo", "SD"),
    ("Speleološki klub", "SK"),
    ("Speleološka udruga", "SU"),
    ("Speleološka sekcija", "SS"),
    ("Hrvatsko planinarsko društvo", "HPD"),
    ("Hrvatski planinarski klub", "HPK"),
    ("Planinarsko društvo Sveučilišta", "PDS"),
    ("Planinarsko društvo", "PD"),
)
_PREFIX_LONG = {short: long for long, short in _TYPE_PREFIXES}
# A section's parent acronym ("SO HPD Mosor"); the compact form drops it.
_PARENT_ACRONYMS = frozenset({"HPD", "PDS", "PD", "PK", "HPK"})
_QUOTES = re.compile(r'["“”„«»]')

# A society list in one cell: "SKOL, SO Sv. Jakov Bitelić" or "SUE; SDV".
_LIST_SEPARATOR = re.compile(r"\s*[,;]\s*")
# A canonical carries ", <Grad>", so up to three comma parts may be one society.
_MAX_PARTS = 3


@dataclass(frozen=True)
class Society:
    name: str
    place: str | None = None
    canonical: str | None = None
    short: str | None = None
    aliases: tuple[str, ...] = ()
    plaque: tuple[str, ...] = ()
    #: times CroSpeleo credits it in the export the list was built from;
    #: 0 = not a CroSpeleo organisation (an HPS-only overlay entry).
    crospeleo_uses: int = 0
    curated: bool = False

    @property
    def in_crospeleo(self) -> bool:
        return self.crospeleo_uses > 0

    @property
    def shortest(self) -> str:
        """The form for tight spaces: ``short`` when there is one, else the name."""
        return self.short or self.name

    def forms(self) -> list[str]:
        """Every curated form, the name first, no repeats."""
        seen: dict[str, None] = {}
        for form in (self.name, self.canonical, self.short, *self.aliases):
            if form:
                seen.setdefault(form, None)
        return list(seen)


@dataclass
class SocietyRegistry:
    societies: list[Society]
    _index: dict[str, Society] = field(default_factory=dict, repr=False)
    #: keys two societies both claim as a curated form — a data error,
    #: reported by ``cavedossier societies check`` and kept out of the index.
    conflicts: dict[str, list[str]] = field(default_factory=dict)

    def __post_init__(self) -> None:
        curated: dict[str, list[Society]] = {}
        derived: dict[str, list[Society]] = {}
        for society in self.societies:
            for form in society.forms():
                for key in _keys(form, society.place):
                    _claim(curated, key, society)
            for form in _derived_forms(society.name):
                for key in _keys(form, society.place):
                    _claim(derived, key, society)

        for key, owners in curated.items():
            if len(owners) == 1:
                self._index[key] = owners[0]
            else:
                self.conflicts[key] = [owner.name for owner in owners]
        for key, owners in derived.items():
            if key not in curated and len(owners) == 1:
                self._index[key] = owners[0]

    def to_json(self) -> dict:
        """The table the dashboard's Udruge page draws and filters."""
        rows = [{
            "name": s.name, "short": s.shortest, "place": s.place,
            "canonical": s.canonical, "aliases": list(s.aliases),
            "plaque": list(s.plaque), "uses": s.crospeleo_uses,
            "curated": s.curated,
            # every spelling that resolves here, for the page's search box
            "search": " ".join(normalize_lookup_key(form) for form in
                               [*s.forms(), *_derived_forms(s.name), s.place or ""]),
        } for s in self.societies]
        rows.sort(key=lambda r: (not r["curated"], -r["uses"], r["name"].casefold()))
        return {"societies": rows, "conflicts": self.conflicts}

    def find(self, text: str | None) -> Society | None:
        """The society ``text`` names, or None. Never guesses."""
        if not text or not text.strip():
            return None
        return self._index.get(normalize_lookup_key(text))

    def split(self, raw: str | None) -> list[tuple[str, Society | None]]:
        """A society list as ``(written part, society or None)`` pairs.

        Commas separate societies, but a canonical has its own comma before
        the seat — so at each position the longest run of up to three parts
        that resolves wins, and an unresolved part stands alone.
        """
        parts = [part for part in _LIST_SEPARATOR.split(raw or "") if part]
        out: list[tuple[str, Society | None]] = []
        i = 0
        while i < len(parts):
            for span in range(min(_MAX_PARTS, len(parts) - i), 0, -1):
                text = ", ".join(parts[i:i + span])
                society = self.find(text)
                if society is not None or span == 1:
                    out.append((text, society))
                    i += span
                    break
        return out


def _claim(table: dict[str, list[Society]], key: str, society: Society) -> None:
    owners = table.setdefault(key, [])
    if society not in owners:
        owners.append(society)


def _keys(form: str, place: str | None) -> list[str]:
    keys = [normalize_lookup_key(form)]
    if place:
        keys.append(normalize_lookup_key(f"{form} {place}"))
    return [key for key in keys if key]


def _derived_forms(name: str) -> list[str]:
    """``SO HPD Mosor`` -> ``SO Mosor``, ``Speleološki odsjek HPD Mosor``,
    ``Speleološki odsjek Mosor``: the spellings a recorder produces without
    anyone having to list them."""
    tokens = name.split()
    if len(tokens) < 2 or tokens[0] not in _PREFIX_LONG:
        return []
    prefix, rest = tokens[0], tokens[1:]
    compact = [token for token in rest if token not in _PARENT_ACRONYMS]
    out = []
    for body in (rest, compact):
        if not body:
            continue
        out.append(" ".join([prefix, *body]))
        out.append(" ".join([_PREFIX_LONG[prefix], *body]))
    return [form for form in out if form != name]


def split_canonical(canonical: str) -> tuple[str, str | None]:
    """``Speleološki odsjek HPD "Mosor", Split`` -> (``SO HPD Mosor``, ``Split``)."""
    head, sep, tail = canonical.rpartition(", ")
    if not sep or head.count("(") != head.count(")"):
        # no seat, or the comma sits inside "(Trebinje, BIH)"
        head, tail = canonical, None
    name = " ".join(_QUOTES.sub(" ", head).split())
    for long, short in _TYPE_PREFIXES:
        if name.casefold().startswith(long.casefold() + " "):
            name = f"{short} {name[len(long) + 1:]}"
            break
    return name, (tail.strip() if tail else None)


# ── loading ───────────────────────────────────────────────────────────
def load(overlay_path: Path = OVERLAY_PATH,
         crospeleo_path: Path = CROSPELEO_PATH) -> SocietyRegistry:
    overlay = json.loads(overlay_path.read_text(encoding="utf-8"))
    try:
        crospeleo = json.loads(crospeleo_path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        crospeleo = {}
    uses = {entry["canonical"]: int(entry.get("uses") or 1)
            for entry in crospeleo.get("organizations", [])}
    uses_by_key = {normalize_lookup_key(c): n for c, n in uses.items()}

    societies: list[Society] = []
    covered: set[str] = set()
    for entry in overlay.get("societies", []):
        canonical = entry.get("canonical")
        derived_name, derived_place = (split_canonical(canonical)
                                       if canonical else (None, None))
        key = normalize_lookup_key(canonical) if canonical else None
        if key:
            covered.add(key)
        societies.append(Society(
            name=entry.get("name") or derived_name or canonical,
            place=entry.get("place") or derived_place,
            canonical=canonical,
            short=entry.get("short"),
            aliases=tuple(entry.get("aliases") or ()),
            plaque=tuple(entry.get("plaque") or ()),
            crospeleo_uses=uses_by_key.get(key, 0) if key else 0,
            curated=True,
        ))

    # Every CroSpeleo organisation the overlay does not mention. CroSpeleo
    # spells a few of them two ways ("SD Meandar" beside the full canonical,
    # a stray double space), so any key the overlay already claims is the
    # overlay's, and two spellings with one key become one entry.
    for society in societies:
        for form in society.forms():
            covered.update(_keys(form, society.place))
    taken = {normalize_lookup_key(s.shortest) for s in societies}
    auto: dict[str, tuple[str, str, str | None, int]] = {}
    for canonical, count in uses.items():
        name, place = split_canonical(canonical)
        keys = {normalize_lookup_key(canonical), *_keys(name, place)}
        if keys & covered:
            continue
        key = normalize_lookup_key(canonical)
        if key in auto:
            first = auto[key]
            auto[key] = (*first[:3], first[3] + count)
        else:
            auto[key] = (canonical, name, place, count)
    auto = list(auto.values())
    rule_shorts = Counter(society_shorthand(name) for _c, name, _p, _n in auto)
    for canonical, name, place, count in auto:
        short = society_shorthand(name)
        if (short is None or rule_shorts[short] > 1
                or normalize_lookup_key(short) in taken):
            short = None    # ambiguous: the name stays the shortest form
        societies.append(Society(name=name, place=place, canonical=canonical,
                                 short=short, crospeleo_uses=count))
    return SocietyRegistry(societies)


@lru_cache(maxsize=1)
def registry() -> SocietyRegistry:
    """The shipped registry, loaded once. Fail-soft: an unreadable overlay
    gives an empty registry and every caller falls back to the rule."""
    try:
        return load()
    except (OSError, ValueError, KeyError):
        return SocietyRegistry([])


def find(text: str | None) -> Society | None:
    return registry().find(text)


# ── fitting a list into a tight space ─────────────────────────────────
def society_ladder(raw: str | None,
                   reg: SocietyRegistry | None = None) -> list[str]:
    """The candidate renderings of a society list, longest first.

    1. as written;
    2. each registered society by its working name (``Speleološka udruga
       "Estavela", Kastav`` -> ``SU Estavela``), where that is shorter;
    3. each society by its short form (``SO Sv. Jakov Bitelić`` ->
       ``SS Sv. JB``); one the registry lacks gets the abbreviation rule, and
       anything unrecognised is kept exactly as written.

    Repeats are dropped, so a list that is already short is one candidate.
    """
    if not raw or not raw.strip():
        return []
    reg = reg if reg is not None else registry()
    pairs = reg.split(raw)
    named = [_shorter(part, society.name) if society else part
             for part, society in pairs]
    shortest = [
        _shorter(current, society.shortest) if society
        else (society_shorthand(part) or part)
        for current, (part, society) in zip(named, pairs)
    ]
    ladder: list[str] = []
    for text in (raw.strip(), ", ".join(named), ", ".join(shortest)):
        if text not in ladder:
            ladder.append(text)
    return ladder


def recognised_count(raw: str | None, reg: SocietyRegistry | None = None) -> int:
    """How many entries of a list are societies (registry or rule)."""
    reg = reg if reg is not None else registry()
    return sum(1 for part, society in reg.split(raw)
               if society is not None or society_shorthand(part))


def fit_societies(raw: str | None, fits: Callable[[str], bool],
                  reg: SocietyRegistry | None = None) -> str | None:
    """The longest rendering of ``raw`` that ``fits``; the shortest if none."""
    ladder = society_ladder(raw, reg)
    if not ladder:
        return None
    for text in ladder:
        if fits(text):
            return text
    return ladder[-1]


def _shorter(current: str, candidate: str) -> str:
    return candidate if len(candidate) < len(current) else current


# ── building the ground truth from a CroSpeleo objects export ────────
# The export lists several organisations in one cell joined by ", " — the
# same separator a canonical uses before its seat. Without CroSpeleo's own
# organisation list to match against, a part is told apart from a seat by
# shape: a seat has no quotes, no organisation word and at most three words.
_ORG_WORD = re.compile(
    r"(?i)dru[sš]tv|klub|club|odsjek|odsek|udrug|sekcij|muzej|stanic|institut"
    r"|zavod|d\.o\.o|savez|ustanov|grotto|grupa|groupe|skupina|speleo|jamar"
    r"|ministarstv|sveu[cč]ili|akademij|policij|uprav|odjel|odio|federation"
    r"|f[ée]d[ée]ration|team|verband|societ|commission|association|[\"„“”']"
)
_EXPORT_COLUMNS = ("Predala udruga", "Istražile udruge", "Izradile udruge (nacrt)",
                   "Dostavile udruge (fotografija ulaza)")


def _is_seat(part: str) -> bool:
    return not _ORG_WORD.search(part) and len(part.split()) <= 3


def split_export_cell(raw: str, known: set[str] | None = None) -> list[str]:
    """One export cell -> its canonical organisation names.

    A lowercase-initial part continues the previous one (``"Osmica" društvo za
    planinarenje, istraživanje i očuvanje …``). Then, greedily, the longest run
    that is a ``known`` canonical wins; otherwise an organisation part takes
    the seat that follows it.
    """
    parts: list[str] = []
    for part in raw.split(", "):
        part = part.strip()
        if parts and part[:1].islower():
            parts[-1] += ", " + part
        elif part:
            parts.append(part)
    known = known or set()
    out: list[str] = []
    i = 0
    while i < len(parts):
        for span in range(len(parts) - i, 1, -1):
            text = ", ".join(parts[i:i + span])
            if text in known:
                out.append(text)
                i += span
                break
        else:
            if i + 1 < len(parts) and _is_seat(parts[i + 1]) and not _is_seat(parts[i]):
                out.append(f"{parts[i]}, {parts[i + 1]}")
                i += 2
            else:
                out.append(parts[i])
                i += 1
    return out


def build_from_export(export_path: Path) -> dict:
    """The ``crospeleo_organizations.json`` payload for one objects export."""
    import openpyxl

    workbook = openpyxl.load_workbook(export_path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        header = [str(cell or "").strip() for cell in next(rows)]
        columns = [header.index(name) for name in _EXPORT_COLUMNS if name in header]
        if not columns:
            raise ValueError(
                f"{export_path.name}: none of the columns {', '.join(_EXPORT_COLUMNS)} "
                "— is this a CroSpeleo objects export?")
        cells = [str(row[i]).strip() for row in rows for i in columns
                 if i < len(row) and row[i]]
    finally:
        workbook.close()

    # A cell with at most one ", " is one organisation as CroSpeleo spells it
    # ("Breganja, Bregana") — those seed the known set for the long cells.
    # Two seatless organisations in one cell look the same, hence the seat test.
    known = {cell for cell in cells
             if ", " not in cell
             or (cell.count(", ") == 1 and _is_seat(cell.split(", ")[1]))}
    counts = Counter(org for cell in cells for org in split_export_cell(cell, known))
    seats = {split_canonical(org)[1] for org in counts} - {None}
    organizations = sorted(
        ({"canonical": org, "uses": n} for org, n in counts.items()
         if not (org in seats and _is_seat(org))),     # a stray bare seat
        key=lambda entry: (-entry["uses"], entry["canonical"]),
    )
    return {
        "_about": "Ground truth for the registar udruga: every organisation "
                  "CroSpeleo credits, with how often. GENERATED by `cavedossier "
                  "societies build` — do not edit; curate societies.json instead.",
        "source": export_path.name,
        "built": date.today().isoformat(),
        "organizations": organizations,
    }


def write_crospeleo(payload: dict, path: Path = CROSPELEO_PATH) -> Path:
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=1) + "\n",
                    encoding="utf-8")
    registry.cache_clear()
    return path
